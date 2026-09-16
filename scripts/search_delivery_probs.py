"""Search Atlas sheets for delivery-outcome probability formulas matching C# service."""
import openpyxl

wb = openpyxl.load_workbook(
    r"c:\Users\b.carson\Downloads\Atlas 196 (2).xlsm",
    data_only=False,
    keep_vba=True,
)

needles = [
    "0.034",
    "0.064",
    "0.051",
    "0.037",
    "0.0023",
    "0.00175",
    "1.23",
    "1.15",
    "FourVolume",
    "1l",
    "1w",
    "5w",
    "7n",
    "1n",
    "2n",
    "5n",
    "delivery",
    "Delivery",
    "extras",
    "no ball",
    "wide",
]

sheets = ["Pricing", "Support", "Support Calcs", "SRL", "Scoring", "Input", "Priority", "UI"]
for sheet in sheets:
    if sheet not in wb.sheetnames:
        print(f"SKIP missing {sheet}")
        continue
    ws = wb[sheet]
    print(f"\n===== {sheet} =====")
    hits = 0
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 2500), max_col=min(ws.max_column or 1, 80)):
        for cell in row:
            v = cell.value
            if v is None:
                continue
            s = str(v)
            if len(s) > 200:
                s_check = s[:200]
            else:
                s_check = s
            if any(n in s_check for n in needles):
                out = s if len(s) <= 160 else s[:157] + "..."
                print(f"{cell.coordinate}: {out}")
                hits += 1
                if hits > 80:
                    print("...truncated...")
                    break
        if hits > 80:
            break
    print(f"hits={hits}")
